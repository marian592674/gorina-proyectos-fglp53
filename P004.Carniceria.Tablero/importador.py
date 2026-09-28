"""Carga inicial y snapshots diarios de Carnisoft. Uso: python importador.py inicial|diario"""
import csv
import datetime as dt
import gzip
import json
import re
import sqlite3
import sys
import shutil
import time
from collections import Counter
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
import db as db_module

ROOT = Path(__file__).resolve().parent
CONFIG = json.loads((ROOT / 'config.json').read_text(encoding='utf-8'))

def path(key):
    p = Path(CONFIG[key])
    resolved = p if p.is_absolute() else ROOT / p
    if not resolved.exists():
        s = str(resolved)
        if s.upper().startswith('J:'):
            unc = Path(r'\\svr\d' + s[2:])
            if unc.exists():
                return unc
    return resolved

def code(value):
    s = str(value or '').strip()
    return str(int(float(s))) if re.fullmatch(r'\d+(?:\.0+)?', s) else s

def decimal_ar(value):
    return float(Decimal(str(value).strip().replace('.', '').replace(',', '.')))

def catalog():
    file = path('catalogo')
    if not file.is_file():
        raise ValueError(f'No se puede leer catálogo: {file}')
    ws = load_workbook(file, read_only=True, data_only=True)['Hoja2']
    result = {}
    for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if row[0] is None:
            continue
        key = code(row[0])
        item = (str(row[2] or '').strip(), str(row[3] or '').strip(),
                str(row[4] or '').strip(), row[5])
        if key in result:
            raise ValueError(f'Código duplicado en catálogo, fila {index}: {key}')
        result[key] = item
    return result

def connect():
    db_module.init_db()
    return db_module.get_db()

INSERT = 'INSERT INTO ventas (fecha, codigo, producto, tipo, comprobante, cliente, kg, precio, subtotal, grupo, agrupa, expone, condicion, neto, kg_neto, tipo_venta, donacion, origen, fila) VALUES (' + ','.join(['?'] * 19) + ')'

def historical(file):
    ws = load_workbook(file, read_only=True, data_only=True)['3. Ventas']
    for i, r in enumerate(ws.iter_rows(min_row=8, values_only=True), 8):
        date = r[2]
        if not isinstance(date, (dt.date, dt.datetime)):
            continue
        yield (date.strftime('%Y-%m-%d'), code(r[3]), str(r[4] or '').strip(),
               str(r[6] or '').strip(), str(r[7] or '').strip(),
               str(r[9] or '').strip(), float(r[10] or 0), float(r[11] or 0),
               float(r[12] or 0), str(r[13] or '').strip(), str(r[14] or '').strip(),
               str(r[15] or '').strip(), float(r[16] or 1), float(r[17] or 0),
               float(r[18] or 0), str(r[19] or '').strip(),
               str(r[20] or '').strip(), file.name, i)

def initial():
    file = path('historico_inicial')
    if not file.is_file():
        raise ValueError(f'Falta el archivo histórico inicial: {file}')
    cat = catalog()
    if file.suffix.lower() == '.csv':
        rows, dates = parse_csv(file, cat, multiple=True)
        if max(dates) >= dt.date.today().isoformat():
            raise ValueError('El histórico inicial incluye el día actual; usar CSV hasta ayer')
    else:
        rows = historical(file)
    with connect() as db:
        if db.execute('SELECT COUNT(*) FROM ventas').fetchone()[0]:
            raise ValueError('La base ya tiene ventas. Carga inicial cancelada para proteger datos.')
        count = 0
        for row in rows:
            db.execute(INSERT, row)
            count += 1
        if count == 0:
            raise ValueError('El histórico no contiene ventas')
        db.execute('INSERT OR REPLACE INTO estado VALUES (?,?)', ('inicial', dt.datetime.now().isoformat(timespec='seconds')))
    print(f'Carga inicial: {count} renglones')

def parse_csv(file, cat, multiple=False):
    rows, dates, missing = [], set(), Counter()
    control = None
    with file.open(encoding='cp1252', newline='') as source:
        for lineno, r in enumerate(csv.reader(source, delimiter=';'), 1):
            if len(r) > 10 and r[7].strip().lower() == 'total general':
                control = (decimal_ar(r[8]), decimal_ar(r[10]))
            if len(r) < 11 or not re.fullmatch(r'\d{2}/\d{2}/\d{4}', r[0].strip()):
                continue
            date = dt.datetime.strptime(r[0].strip(), '%d/%m/%Y').date().isoformat()
            dates.add(date)
            key = code(r[1])
            entry = cat.get(key)
            if not entry or any(not x or x == '#N/A' for x in entry[:3]) or not isinstance(entry[3], (float, int)):
                missing[key] += 1
                continue
            agrupa, expone, grupo, iva = entry
            gross, kilos = decimal_ar(r[10]), decimal_ar(r[8])
            rows.append((date, key, r[2].strip(), r[4].strip(), r[5].strip(),
                         r[7].strip(), kilos, decimal_ar(r[9]), gross,
                         grupo, agrupa, expone, 1 + iva, gross / (1 + iva),
                         round(kilos, 2), r[5].strip()[:1], '', file.name, lineno))
    if not dates or not rows or (len(dates) != 1 and not multiple):
        raise ValueError(f'CSV vacío o con fechas mezcladas: {file.name}')
    if missing:
        raise ValueError(f'Códigos sin clasificación/IVA en {file.name}: {dict(missing)}')
    if control is None or abs(sum(r[6] for r in rows) - control[0]) > .01:
        raise ValueError(f'El total de kilos no coincide con Total General: {file.name}')
    difference = sum(r[8] for r in rows) - control[1]
    tol_hist = float(CONFIG.get('tolerancia_historico_diferencia', 5000000000.0))
    if abs(difference) > .01:
        if multiple and abs(difference - tol_hist) < .01:
            print(f'ADVERTENCIA histórico: suma de renglones ${tol_hist:,.2f} mayor que Total General impreso. Se usan los renglones individuales; revisar con Carnisoft.')
        else:
            raise ValueError(f'El total de importe no coincide con Total General: {file.name}; diferencia={difference:.2f}')
    return rows, dates if multiple else dates.pop()

def backup(db):
    folder = path('respaldo')
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / ('carniceria_' + dt.date.today().isoformat() + '.sqlite.gz')
    temporary = folder / '_respaldo_temporal.sqlite'
    dest = sqlite3.connect(temporary)
    try:
        db.backup(dest)
    finally:
        dest.close()
    with temporary.open('rb') as source, gzip.open(target, 'wb', compresslevel=6) as compressed:
        shutil.copyfileobj(source, compressed)
    if temporary.is_file():
        temporary.unlink()
    cutoff = dt.date.today() - dt.timedelta(days=14)
    for old in folder.glob('carniceria_*.sqlite.gz'):
        try:
            date = dt.date.fromisoformat(old.name[11:21])
            if date < cutoff:
                old.unlink()
        except ValueError:
            pass
    print('Respaldo local:', target)

def daily():
    directory = path('carpeta_carnisoft')
    if not directory.is_dir():
        raise ValueError(f'No se puede leer carpeta: {directory}. Usar ruta UNC si J: no existe para la tarea programada.')
    cat = catalog()
    latest = {}
    for f in directory.glob('*.csv'):
        match = re.fullmatch(r'(\d{8})_(\d{6})\.csv', f.name, re.I)
        if match:
            latest[match.group(1)] = max(latest.get(match.group(1), f), f, key=lambda p: p.name)
    if not latest:
        raise ValueError('No hay CSV con nombre AAAAMMDD_HHMMSS.csv en la carpeta')
    processed = 0
    with connect() as db:
        for stamp, file in sorted(latest.items()):
            date = dt.datetime.strptime(stamp, '%Y%m%d').date()
            if date > dt.date.today():
                continue
            if date == dt.date.today() and dt.datetime.now().hour < 21:
                continue
            if date == dt.date.today() and int(file.stem.split('_')[1]) < 210000:
                raise ValueError(f'Falta descarga de cierre (21:00 o posterior) para {date}; último archivo: {file.name}')
            if time.time() - file.stat().st_mtime < 60:
                continue
            old = db.execute('SELECT archivo,filas FROM cargas WHERE fecha=?', (date.isoformat(),)).fetchone()
            if old and old[0] == file.name:
                continue
            rows, actual = parse_csv(file, cat)
            if actual != date.isoformat():
                raise ValueError(f'Fecha de archivo y ventas diferente: {file.name}')
            if old and len(rows) < old[1]:
                raise ValueError(f'Nueva descarga más corta que la anterior: {file.name}')
            with db:
                db.execute('DELETE FROM ventas WHERE fecha=?', (actual,))
                db.executemany(INSERT, rows)
                db.execute('INSERT OR REPLACE INTO cargas VALUES (?,?,?,?)',
                           (actual, file.name, len(rows), dt.datetime.now().isoformat(timespec='seconds')))
            print(f'{actual}: {len(rows)} renglones desde {file.name}')
            processed += 1
        if processed:
            backup(db)
    print(f'Días actualizados: {processed}')

if __name__ == '__main__':
    try:
        if sys.argv[1:] == ['inicial']:
            initial()
        elif sys.argv[1:] == ['diario']:
            daily()
        else:
            raise ValueError('Uso: python importador.py inicial | diario')
    except Exception as exc:
        print('ERROR:', exc, file=sys.stderr)
        sys.exit(1)
